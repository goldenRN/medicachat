from __future__ import annotations

import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from backend.config import UPLOAD_DIR
from backend.store import ensure_bootstrap, replace_all_employees
from backend.text_utils import sanitize_filename, utc_now


WORKBOOK_DEFAULT = ROOT_DIR / "data" / "uploads" / "Ажилчдын-мэдээлэл" / "1779941830204-Emp.info-2026.05.xlsx"

SHEET_CONFIGS = {
    "Clinic-Удирдлага, захиргаа": {
        "categoryKey": "clinic_management_admin",
        "categoryNameMn": "Clinic удирдлага, захиргааны ажилчид",
        "categoryNameEn": "Clinic Management and Administration",
    },
    "Clinic-Эмнэлгийн чиг үүргийн": {
        "categoryKey": "clinic_medical",
        "categoryNameMn": "Clinic эмнэлгийн чиг үүргийн ажилчид",
        "categoryNameEn": "Clinic Medical Function",
    },
    "OT Staff": {
        "categoryKey": "ot_staff",
        "categoryNameMn": "OT Staff / Оюу Толгойн staff",
        "categoryNameEn": "OT Staff",
    },
    "Clinic-Бусад чиг үүргийн": {
        "categoryKey": "clinic_other",
        "categoryNameMn": "Clinic бусад чиг үүргийн ажилчид",
        "categoryNameEn": "Clinic Other Function",
    },
}


def is_placeholder(value: object) -> bool:
    compact = str(value or "").strip()
    if not compact:
        return True
    return compact.lower() in {"*", "n/a", "na", "none", "null", "байхгүй"}


def normalize_value(value: object, *, treat_as_date: bool = False) -> str:
    if value is None:
        return ""
    if treat_as_date and isinstance(value, (int, float)):
        try:
            return from_excel(value).strftime("%Y-%m-%d")
        except Exception:
            pass
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def detect_text_language(value: object) -> str:
    compact = str(value or "")
    latin_count = len(re.findall(r"[A-Za-z]", compact))
    cyrillic_count = len(re.findall(r"[А-Яа-яӨөҮүЁё]", compact))
    if cyrillic_count > latin_count:
        return "mn"
    if latin_count > cyrillic_count:
        return "en"
    return ""


def find_header_row(worksheet) -> int:
    for row_index in range(1, min(worksheet.max_row, 12) + 1):
        values = [str(worksheet.cell(row_index, column).value or "") for column in range(1, worksheet.max_column + 1)]
        if any("Last name" in value for value in values):
            return row_index
    raise ValueError(f"Header row not found for sheet {worksheet.title}")


def build_header_map(worksheet, header_row: int) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for column_index in range(1, worksheet.max_column + 1):
        header_value = str(worksheet.cell(header_row, column_index).value or "").strip().lower()
        if not header_value:
            continue
        if "last name" in header_value:
            header_map["last_name"] = column_index
        elif "first name" in header_value:
            header_map["first_name"] = column_index
        elif "photo" in header_value:
            header_map["photo"] = column_index
        elif "position" in header_value:
            header_map["position"] = column_index
        elif "register" in header_value or "регист" in header_value:
            header_map["register_number"] = column_index
        elif "email" in header_value:
            header_map["email"] = column_index
        elif "telephone" in header_value or " утас" in header_value:
            header_map["phone"] = column_index
        elif "duty phone" in header_value:
            header_map["duty_phone"] = column_index
        elif "dob" in header_value:
            header_map["date_of_birth"] = column_index
        elif "date of hire" in header_value or "ажилд орсон" in header_value:
            header_map["hire_date"] = column_index
        elif "home address" in header_value or "гэрийн хаяг" in header_value:
            header_map["home_address"] = column_index
        elif "notes" in header_value or "тэмдэглэл" in header_value:
            header_map["notes"] = column_index
        elif header_value not in {"№", "za"} and "column" in header_value:
            header_map["extra_info"] = column_index
    return header_map


def collect_sheet_images(worksheet, photo_column: int) -> dict[int, tuple[bytes, str]]:
    images_by_row: dict[int, tuple[bytes, str]] = {}
    for image in getattr(worksheet, "_images", []):
        anchor = image.anchor._from
        row_index = int(anchor.row) + 1
        column_index = int(anchor.col) + 1
        if row_index < 2 or column_index != photo_column:
            continue
        image_format = getattr(image, "format", None) or "png"
        images_by_row[row_index] = (image._data(), str(image_format).lower())
    return images_by_row


def assign_language_field(record: dict[str, str], base_key: str, value: object) -> None:
    normalized = normalize_value(value)
    if is_placeholder(normalized):
        return
    language = detect_text_language(normalized)
    if language == "mn":
        target_key = f"{base_key}Mn"
    elif language == "en":
        target_key = f"{base_key}En"
    else:
        target_key = f"{base_key}Mn" if not record.get(f"{base_key}Mn") else f"{base_key}En"
    if not record.get(target_key):
        record[target_key] = normalized


def unique_values(values: list[str]) -> list[str]:
    seen: list[str] = []
    for value in values:
        compact = normalize_value(value)
        if is_placeholder(compact) or compact in seen:
            continue
        seen.append(compact)
    return seen


def select_best_value(values: list[str]) -> str:
    options = unique_values(values)
    if not options:
        return ""
    options.sort(key=lambda item: (len(item), item), reverse=True)
    return options[0]


def clean_photo_dir(photo_dir: Path) -> None:
    if photo_dir.exists():
        shutil.rmtree(photo_dir)
    photo_dir.mkdir(parents=True, exist_ok=True)


def parse_workbook(workbook_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path)
    photo_dir = UPLOAD_DIR / "employee-photos"
    clean_photo_dir(photo_dir)
    records: list[dict[str, str]] = []

    for worksheet in workbook.worksheets:
        config = SHEET_CONFIGS.get(worksheet.title)
        if not config:
            continue

        header_row = find_header_row(worksheet)
        header_map = build_header_map(worksheet, header_row)
        photo_column = header_map.get("photo", 4)
        images_by_row = collect_sheet_images(worksheet, photo_column)

        row_index = header_row + 1
        employee_count = 0
        while row_index <= worksheet.max_row:
            first_name = worksheet.cell(row_index, header_map.get("first_name", 3)).value
            start_marker = worksheet.cell(row_index, 1).value
            if first_name and start_marker is not None:
                employee_count += 1
                primary_row = [worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)]
                secondary_row = None
                if row_index + 1 <= worksheet.max_row:
                    next_first_name = worksheet.cell(row_index + 1, header_map.get("first_name", 3)).value
                    next_start_marker = worksheet.cell(row_index + 1, 1).value
                    if next_first_name and next_start_marker in {None, ""}:
                        secondary_row = [worksheet.cell(row_index + 1, column).value for column in range(1, worksheet.max_column + 1)]

                record = {
                    "id": f"{config['categoryKey']}-{employee_count:03d}",
                    "categoryKey": config["categoryKey"],
                    "categoryNameMn": config["categoryNameMn"],
                    "categoryNameEn": config["categoryNameEn"],
                    "sourceSheet": worksheet.title,
                    "employeeNumber": normalize_value(start_marker),
                    "sortOrder": str(employee_count),
                    "lastNameEn": "",
                    "firstNameEn": "",
                    "lastNameMn": "",
                    "firstNameMn": "",
                    "positionEn": "",
                    "positionMn": "",
                    "registerNumber": "",
                    "emailPrimary": "",
                    "emailSecondary": "",
                    "phonePrimary": "",
                    "phoneSecondary": "",
                    "dutyPhone": "",
                    "dateOfBirth": "",
                    "hireDate": "",
                    "homeAddress": "",
                    "notes": "",
                    "extraInfo": "",
                    "photoStoragePath": "",
                    "createdAt": utc_now(),
                    "updatedAt": utc_now(),
                }

                row_candidates = [primary_row]
                if secondary_row:
                    row_candidates.append(secondary_row)

                for row_values in row_candidates:
                    assign_language_field(record, "lastName", row_values[header_map.get("last_name", 2) - 1])
                    assign_language_field(record, "firstName", row_values[header_map.get("first_name", 3) - 1])
                    assign_language_field(record, "position", row_values[header_map.get("position", 5) - 1])

                record["registerNumber"] = select_best_value([
                    normalize_value(row_values[header_map["register_number"] - 1]) for row_values in row_candidates if "register_number" in header_map
                ])

                emails = unique_values([
                    normalize_value(row_values[header_map["email"] - 1]) for row_values in row_candidates if "email" in header_map
                ])
                record["emailPrimary"] = emails[0] if emails else ""
                record["emailSecondary"] = emails[1] if len(emails) > 1 else ""

                phones = unique_values([
                    normalize_value(row_values[header_map["phone"] - 1]) for row_values in row_candidates if "phone" in header_map
                ])
                record["phonePrimary"] = phones[0] if phones else ""
                record["phoneSecondary"] = phones[1] if len(phones) > 1 else ""

                record["dutyPhone"] = select_best_value([
                    normalize_value(row_values[header_map["duty_phone"] - 1]) for row_values in row_candidates if "duty_phone" in header_map
                ])

                record["dateOfBirth"] = select_best_value([
                    normalize_value(
                        row_values[header_map["date_of_birth"] - 1],
                        treat_as_date=True,
                    )
                    for row_values in row_candidates if "date_of_birth" in header_map
                ])
                record["hireDate"] = select_best_value([
                    normalize_value(
                        row_values[header_map["hire_date"] - 1],
                        treat_as_date=True,
                    )
                    for row_values in row_candidates if "hire_date" in header_map
                ])
                record["homeAddress"] = select_best_value([
                    normalize_value(row_values[header_map["home_address"] - 1]) for row_values in row_candidates if "home_address" in header_map
                ])

                note_values = unique_values([
                    normalize_value(row_values[header_map["notes"] - 1]) for row_values in row_candidates if "notes" in header_map
                ])
                record["notes"] = " | ".join(note_values)

                extra_values = unique_values([
                    normalize_value(row_values[header_map["extra_info"] - 1]) for row_values in row_candidates if "extra_info" in header_map
                ])
                record["extraInfo"] = " | ".join(extra_values)

                image_bytes = None
                image_format = "png"
                for image_row in (row_index, row_index + 1):
                    if image_row in images_by_row:
                        image_bytes, image_format = images_by_row[image_row]
                        break
                if image_bytes:
                    safe_name = sanitize_filename(
                        f"{record['id']}-{record['firstNameEn'] or record['firstNameMn']}-{record['lastNameEn'] or record['lastNameMn']}"
                    ).strip("-") or record["id"]
                    photo_path = photo_dir / f"{safe_name}.{image_format}"
                    photo_path.write_bytes(image_bytes)
                    record["photoStoragePath"] = str(photo_path.relative_to(UPLOAD_DIR))

                records.append(record)
            row_index += 1

    return records


def main() -> None:
    workbook_path = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else WORKBOOK_DEFAULT.resolve()
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    ensure_bootstrap(run_maintenance=False)
    records = parse_workbook(workbook_path)
    replace_all_employees(records)
    print(f"Imported {len(records)} employees from {workbook_path.name}")


if __name__ == "__main__":
    main()
